"""
inventory_engine/routes.py
Flask Blueprint routes for Inventory Intelligence:
- GET /inventory : Standalone Executive UI
- POST /api/inventory/upload : Upload Excel files and create snapshot
- GET /api/inventory/snapshots : List historical snapshots
- GET /api/inventory/analysis/<id> : Get 4-engine analytical insight
- GET /api/inventory/export/<id> : Export professional multi-sheet Excel work order
"""
from __future__ import annotations

import io
from datetime import datetime
from typing import Any, Dict

from flask import Blueprint, jsonify, render_template, request, send_file
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from inventory_engine.analytics import run_inventory_engine
from inventory_engine.parser import parse_stok_gudang, parse_stok_konter
from inventory_engine.repository import (
    get_latest_snapshot_id,
    list_snapshots,
    save_snapshot,
)

inventory_bp = Blueprint("inventory", __name__, template_folder="../templates")


@inventory_bp.route("/inventory", methods=["GET"])
def inventory_dashboard():
    """Renders the standalone Smart Inventory UI."""
    return render_template("inventory.html")


@inventory_bp.route("/api/inventory/snapshots", methods=["GET"])
def get_snapshots_list():
    """Returns JSON list of all snapshots."""
    try:
        snapshots = list_snapshots()
        latest_id = get_latest_snapshot_id()
        return jsonify({
            "success": True,
            "snapshots": snapshots,
            "latest_id": latest_id,
        })
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500


@inventory_bp.route("/api/inventory/upload", methods=["POST"])
def upload_inventory_files():
    """
    Accepts 1 or 2 Excel files:
    - 'file_gudang' (optional if only konter uploaded)
    - 'file_konter' (optional if only gudang uploaded)
    """
    try:
        file_gudang = request.files.get("file_gudang")
        file_konter = request.files.get("file_konter")

        if not file_gudang and not file_konter:
            return jsonify({
                "success": False,
                "error": "Mohon sertakan setidaknya salah satu file: Stok Gudang atau Stok Konter.",
            }), 400

        gudang_items = []
        konter_items = []
        name_g = None
        name_k = None

        if file_gudang and file_gudang.filename:
            name_g = file_gudang.filename
            gudang_items = parse_stok_gudang(file_gudang.read())

        if file_konter and file_konter.filename:
            name_k = file_konter.filename
            k_items, fallback_g = parse_stok_konter(file_konter.read())
            konter_items = k_items
            # If gudang was not uploaded separately, use fallback warehouse stock from konter file
            if not gudang_items and fallback_g:
                gudang_items = fallback_g

        if not gudang_items and not konter_items:
            return jsonify({
                "success": False,
                "error": "File yang diunggah tidak memiliki data baris yang valid.",
            }), 400

        snapshot_id = save_snapshot(
            filename_gudang=name_g,
            filename_konter=name_k,
            gudang_items=gudang_items,
            konter_items=konter_items,
        )

        return jsonify({
            "success": True,
            "message": f"Snapshot #{snapshot_id} berhasil diproses dan disimpan.",
            "snapshot_id": snapshot_id,
            "gudang_sku_count": len(gudang_items),
            "konter_row_count": len(konter_items),
        })
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500


@inventory_bp.route("/api/inventory/analysis/<int:snapshot_id>", methods=["GET"])
def get_analysis(snapshot_id: int):
    """Runs the 4 analytics engines for a given snapshot."""
    try:
        safety_days = int(request.args.get("safety_days", 7))
        result = run_inventory_engine(snapshot_id, safety_days=safety_days)
        return jsonify({
            "success": True,
            "data": result,
        })
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500


@inventory_bp.route("/api/inventory/export/<int:snapshot_id>", methods=["GET"])
def export_excel_work_order(snapshot_id: int):
    """
    Generates an executive multi-sheet Excel work order for logistics, store managers, and buyers.
    """
    try:
        safety_days = int(request.args.get("safety_days", 7))
        data = run_inventory_engine(snapshot_id, safety_days=safety_days)

        wb = openpyxl.Workbook()
        # Remove default sheet
        wb.remove(wb.active)

        # Style helpers
        font_header = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        fill_dark = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        fill_blue = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
        fill_amber = PatternFill(start_color="D97706", end_color="D97706", fill_type="solid")
        fill_red = PatternFill(start_color="DC2626", end_color="DC2626", fill_type="solid")
        fill_purple = PatternFill(start_color="7C3AED", end_color="7C3AED", fill_type="solid")
        thin_border = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1"),
        )
        align_left = Alignment(horizontal="left", vertical="center")
        align_right = Alignment(horizontal="right", vertical="center")
        align_center = Alignment(horizontal="center", vertical="center")

        def style_sheet_headers(ws, title_text, fill_color, headers):
            ws.views.sheetView[0].showGridLines = True
            # Title banner
            ws.merge_cells("A1:G1")
            ws["A1"] = title_text
            ws["A1"].font = Font(name="Arial", size=14, bold=True, color="1E293B")
            ws["A1"].alignment = align_left

            ws["A2"] = f"Digenerate pada: {datetime.now().strftime('%d %B %Y %H:%M:%S')} | Safety Stock Buffer: {safety_days} Hari"
            ws["A2"].font = Font(name="Arial", size=9, italic=True, color="64748B")

            # Column headers at Row 4
            for col_idx, h in enumerate(headers, start=1):
                cell = ws.cell(row=4, column=col_idx, value=h)
                cell.font = font_header
                cell.fill = fill_color
                cell.alignment = align_center
                cell.border = thin_border
            ws.row_dimensions[4].height = 26

        def autofit_cols(ws):
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    if cell.row in (1, 2):
                        continue
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 48)

        # ----------------------------------------------------
        # SHEET 1: STORE PRIORITY MATRIX
        # ----------------------------------------------------
        ws1 = wb.create_sheet(title="Prioritas Konter")
        headers1 = ["Rank", "Kode Konter", "Nama Konter / Lokasi", "Area", "Peringkat Traffic", "Skor Prioritas", "Avg Unit / Hari", "Avg Omset / Hari", "Total Unit Terjual"]
        style_sheet_headers(ws1, "ANCOL STORE PRIORITY MATRIX (RANKING KECEPATAN KONTER)", fill_dark, headers1)

        for row_idx, s in enumerate(data["store_priority_matrix"], start=5):
            ws1.cell(row=row_idx, column=1, value=s["rank"]).alignment = align_center
            ws1.cell(row=row_idx, column=2, value=s["outlet_code"]).alignment = align_center
            ws1.cell(row=row_idx, column=3, value=s["outlet_name"]).alignment = align_left
            ws1.cell(row=row_idx, column=4, value=s["area"]).alignment = align_center
            ws1.cell(row=row_idx, column=5, value=s["tier_badge"]).alignment = align_center
            ws1.cell(row=row_idx, column=6, value=s["priority_score"]).alignment = align_center
            ws1.cell(row=row_idx, column=7, value=s["avg_daily_units"]).alignment = align_right
            ws1.cell(row=row_idx, column=8, value=s["avg_daily_revenue"]).alignment = align_right
            ws1.cell(row=row_idx, column=9, value=s["total_units_sold"]).alignment = align_right
            for c in range(1, 10):
                ws1.cell(row=row_idx, column=c).border = thin_border
            ws1.row_dimensions[row_idx].height = 20
        autofit_cols(ws1)

        # ----------------------------------------------------
        # SHEET 2: RESTOCK RECOMMENDATIONS (GUDANG -> KONTER)
        # ----------------------------------------------------
        ws2 = wb.create_sheet(title="SPK Kirim Gudang")
        headers2 = ["Kode SKU", "Nama Barang", "Supplier", "Konter Tujuan", "Rekomendasi Kirim (Pcs)", "Stok Konter Saat Ini", "Kecepatan Jual (Pcs/Hr)", "Target Stok Aman", "Kekurangan (Defisit)", "Sisa Stok Gudang", "Status Alokasi"]
        style_sheet_headers(ws2, "SURAT PERINTAH KERJA (SPK): RESTOCK GUDANG UTAMA KE KONTER", fill_blue, headers2)

        for row_idx, r in enumerate(data["restock_actionable"], start=5):
            ws2.cell(row=row_idx, column=1, value=r["sku_code"]).alignment = align_center
            ws2.cell(row=row_idx, column=2, value=r["nama_barang"]).alignment = align_left
            ws2.cell(row=row_idx, column=3, value=r["supplier"]).alignment = align_left
            ws2.cell(row=row_idx, column=4, value=f"{r['outlet_code']} - {r['outlet_name']}").alignment = align_left
            ws2.cell(row=row_idx, column=5, value=r["recommended_send_qty"]).alignment = align_right
            ws2.cell(row=row_idx, column=6, value=r["current_stock"]).alignment = align_right
            ws2.cell(row=row_idx, column=7, value=r["daily_velocity"]).alignment = align_right
            ws2.cell(row=row_idx, column=8, value=r["target_stock"]).alignment = align_right
            ws2.cell(row=row_idx, column=9, value=r["deficit"]).alignment = align_right
            ws2.cell(row=row_idx, column=10, value=r["gudang_stock_left"]).alignment = align_right
            ws2.cell(row=row_idx, column=11, value=r["status"]).alignment = align_left
            for c in range(1, 12):
                ws2.cell(row=row_idx, column=c).border = thin_border
            ws2.row_dimensions[row_idx].height = 20
        autofit_cols(ws2)

        # ----------------------------------------------------
        # SHEET 3: MUTASI ANTAR KONTER (CROSS-TRANSFER)
        # ----------------------------------------------------
        ws3 = wb.create_sheet(title="Mutasi Antar Konter")
        headers3 = ["Kode SKU", "Nama Barang", "Supplier", "Konter Asal (Overstock)", "Stok Asal", "Konter Tujuan (Krisis)", "Stok Tujuan", "Qty Mutasi (Pcs)", "Alasan Mutasi & Rekomendasi"]
        style_sheet_headers(ws3, "REKOMENDASI MUTASI STOK ANTAR KONTER (SMART BALANCING)", fill_amber, headers3)

        for row_idx, c_tr in enumerate(data["cross_transfers"], start=5):
            ws3.cell(row=row_idx, column=1, value=c_tr["sku_code"]).alignment = align_center
            ws3.cell(row=row_idx, column=2, value=c_tr["nama_barang"]).alignment = align_left
            ws3.cell(row=row_idx, column=3, value=c_tr["supplier"]).alignment = align_left
            ws3.cell(row=row_idx, column=4, value=f"{c_tr['from_outlet_code']} - {c_tr['from_outlet_name']}").alignment = align_left
            ws3.cell(row=row_idx, column=5, value=c_tr["from_current_stock"]).alignment = align_right
            ws3.cell(row=row_idx, column=6, value=f"{c_tr['to_outlet_code']} - {c_tr['to_outlet_name']}").alignment = align_left
            ws3.cell(row=row_idx, column=7, value=c_tr["to_current_stock"]).alignment = align_right
            ws3.cell(row=row_idx, column=8, value=c_tr["transfer_qty"]).alignment = align_right
            ws3.cell(row=row_idx, column=9, value=c_tr["reason"]).alignment = align_left
            for c in range(1, 10):
                ws3.cell(row=row_idx, column=c).border = thin_border
            ws3.row_dimensions[row_idx].height = 20
        autofit_cols(ws3)

        # ----------------------------------------------------
        # SHEET 4: RADAR STOK KRITIS & HABIS TOTAL
        # ----------------------------------------------------
        ws4 = wb.create_sheet(title="Kritis & Habis")
        headers4 = ["Status Alert", "Kode SKU", "Nama Barang", "Supplier", "Konter", "Stok Saat Ini", "Kecepatan Jual (Pcs/Hr)", "Estimasi Runway", "Potensi Omset Hilang / Hari"]
        style_sheet_headers(ws4, "RADAR STOK KRITIS & HABIS TOTAL DI KONTER AKTIF", fill_red, headers4)

        combined_alerts = [
            ("HABIS TOTAL", a) for a in data["stockout_alerts"]
        ] + [
            ("KRITIS (<48 JAM)", a) for a in data["critical_alerts"]
        ]

        for row_idx, (status, a) in enumerate(combined_alerts, start=5):
            ws4.cell(row=row_idx, column=1, value=status).alignment = align_center
            ws4.cell(row=row_idx, column=2, value=a["sku_code"]).alignment = align_center
            ws4.cell(row=row_idx, column=3, value=a["nama_barang"]).alignment = align_left
            ws4.cell(row=row_idx, column=4, value=a["supplier"]).alignment = align_left
            ws4.cell(row=row_idx, column=5, value=f"{a['outlet_code']} - {a['outlet_name']}").alignment = align_left
            ws4.cell(row=row_idx, column=6, value=a["stok_qty"]).alignment = align_right
            ws4.cell(row=row_idx, column=7, value=a["daily_velocity"]).alignment = align_right
            runway = f"{a.get('runway_hours', 0)} Jam" if status.startswith("KRITIS") else "HABIS (0 Hari)"
            ws4.cell(row=row_idx, column=8, value=runway).alignment = align_center
            ws4.cell(row=row_idx, column=9, value=a.get("lost_daily_sales_est", 0)).alignment = align_right
            for c in range(1, 10):
                ws4.cell(row=row_idx, column=c).border = thin_border
            ws4.row_dimensions[row_idx].height = 20
        autofit_cols(ws4)

        # ----------------------------------------------------
        # SHEET 5: SARAN BARANG OVERSTOCK & TINDAKAN
        # ----------------------------------------------------
        ws5 = wb.create_sheet(title="Saran Overstock")
        headers5 = ["Kode SKU", "Nama Barang", "Supplier", "Kategori", "Total Lebihan Stok", "Konter Terdampak", "Rekomendasi Strategi Retail", "Detail Instruksi Operasional"]
        style_sheet_headers(ws5, "PANDUAN STRATEGI PENANGANAN BARANG OVERSTOCK & SLOW MOVING", fill_purple, headers5)

        for row_idx, o in enumerate(data["overstock_advisory"], start=5):
            ws5.cell(row=row_idx, column=1, value=o["sku_code"]).alignment = align_center
            ws5.cell(row=row_idx, column=2, value=o["nama_barang"]).alignment = align_left
            ws5.cell(row=row_idx, column=3, value=o["supplier"]).alignment = align_left
            ws5.cell(row=row_idx, column=4, value=o["kategori"]).alignment = align_center
            ws5.cell(row=row_idx, column=5, value=o["total_excess_units"]).alignment = align_right
            ws5.cell(row=row_idx, column=6, value=o["affected_outlets"]).alignment = align_left
            ws5.cell(row=row_idx, column=7, value=o["action_title"]).alignment = align_left
            ws5.cell(row=row_idx, column=8, value=o["action_detail"]).alignment = align_left
            for c in range(1, 9):
                ws5.cell(row=row_idx, column=c).border = thin_border
            ws5.row_dimensions[row_idx].height = 20
        autofit_cols(ws5)

        output_io = io.BytesIO()
        wb.save(output_io)
        output_io.seek(0)
        export_filename = f"SPK_INVENTORY_INTELLIGENCE_ANCOL_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"

        return send_file(
            output_io,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=export_filename,
        )
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500

